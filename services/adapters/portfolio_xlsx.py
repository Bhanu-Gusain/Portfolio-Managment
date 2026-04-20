"""Adapter for the consolidated multi-sheet Groww portfolio xlsx.

Single workbook with four sheets:
    - MF Holdings          (summary rows on top, real header on row index 2)
    - MF Transactions      (clean header on row 0)
    - Stocks Holdings      (summary rows on top, real header on row index 5)
    - Stocks Transactions  (clean header on row 0)

Unlike the other adapters, one workbook yields rows for *four* canonical
targets, so `parse()` returns a `dict[target, DataFrame]`. The router in
`services/adapters/__init__.py` unpacks that into multiple (target, df) pairs.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

NAME = "portfolio_xlsx"
TARGET = "multi"

REQUIRED_SHEETS = {
    "MF Holdings",
    "MF Transactions",
    "Stocks Holdings",
    "Stocks Transactions",
}


def matches(path: Path) -> bool:
    if path.suffix.lower() not in {".xlsx", ".xls"}:
        return False
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = set(wb.sheetnames)
        wb.close()
    except Exception:
        return False
    return REQUIRED_SHEETS.issubset(sheets)


def parse(path: Path) -> dict[str, pd.DataFrame]:
    xl = pd.ExcelFile(path, engine="openpyxl")

    raw_stock_tx = pd.read_excel(xl, sheet_name="Stocks Transactions")
    name_to_symbol = _build_symbol_map(raw_stock_tx)

    return {
        "stock_holdings": _parse_stock_holdings(xl, name_to_symbol),
        "mf_holdings": _parse_mf_holdings(xl),
        "stock_transactions": _parse_stock_transactions(raw_stock_tx),
        "mf_transactions": _parse_mf_transactions(xl),
    }


# ---------- sheet parsers ----------

def _parse_stock_holdings(xl: pd.ExcelFile, name_to_symbol: dict[str, str]) -> pd.DataFrame:
    raw = pd.read_excel(xl, sheet_name="Stocks Holdings", header=None)
    header_row = _find_header_row(raw, must_contain={"Stock Name", "Quantity", "Average buy price"})
    df = pd.read_excel(xl, sheet_name="Stocks Holdings", header=header_row)
    df = df.dropna(how="all").dropna(subset=["Stock Name"])

    names = df["Stock Name"].astype(str).str.strip()
    symbols = names.map(lambda n: name_to_symbol.get(n, "")).astype(str)

    out = pd.DataFrame({
        "symbol": symbols.str.upper(),
        "exchange": "NSE",
        "quantity": pd.to_numeric(df["Quantity"], errors="coerce"),
        "avg_cost": pd.to_numeric(df["Average buy price"], errors="coerce"),
        "currency": "INR",
    })
    return out.dropna(subset=["quantity", "avg_cost"]).reset_index(drop=True)


def _parse_mf_holdings(xl: pd.ExcelFile) -> pd.DataFrame:
    raw = pd.read_excel(xl, sheet_name="MF Holdings", header=None)
    header_row = _find_header_row(raw, must_contain={"Scheme Name", "Units", "Invested Value"})
    df = pd.read_excel(xl, sheet_name="MF Holdings", header=header_row)
    df = df.dropna(how="all").dropna(subset=["Scheme Name"])

    units = pd.to_numeric(df["Units"], errors="coerce")
    invested = pd.to_numeric(df["Invested Value"], errors="coerce")
    avg_nav = invested / units

    out = pd.DataFrame({
        "isin": "",
        "scheme_name": df["Scheme Name"].astype(str).str.strip(),
        "units": units,
        "avg_nav": avg_nav,
    })
    return out.dropna(subset=["units", "avg_nav"]).reset_index(drop=True)


def _parse_stock_transactions(df: pd.DataFrame) -> pd.DataFrame:
    df = df.dropna(how="all").dropna(subset=["Symbol"])

    qty = pd.to_numeric(df["Quantity"], errors="coerce")
    value = _numeric_with_commas(df["Value"])
    price = value / qty

    out = pd.DataFrame({
        "date": pd.to_datetime(df["Execution date and time"], dayfirst=True, errors="coerce"),
        "symbol": df["Symbol"].astype(str).str.strip().str.upper(),
        "exchange": "NSE",
        "side": df["Type"].astype(str).str.lower().str.strip(),
        "quantity": qty,
        "price": price,
        "charges": 0.0,
        "notes": "",
    })
    out["side"] = out["side"].replace({"buy": "buy", "sell": "sell", "purchase": "buy", "redemption": "sell"})
    out["date"] = out["date"].dt.strftime("%Y-%m-%d")
    return out.dropna(subset=["date", "symbol", "quantity", "price"]).reset_index(drop=True)


def _parse_mf_transactions(xl: pd.ExcelFile) -> pd.DataFrame:
    df = pd.read_excel(xl, sheet_name="MF Transactions")
    df = df.dropna(how="all").dropna(subset=["Scheme Name"])

    out = pd.DataFrame({
        "date": pd.to_datetime(df["Date"], errors="coerce", dayfirst=True),
        "isin": "",
        "scheme_name": df["Scheme Name"].astype(str).str.strip(),
        "side": df["Transaction Type"].astype(str).str.lower().str.strip(),
        "units": pd.to_numeric(df["Units"], errors="coerce"),
        "nav": pd.to_numeric(df["NAV"], errors="coerce"),
        "amount": _numeric_with_commas(df["Amount"]),
        "folio": "",
        "notes": "",
    })
    out["side"] = out["side"].replace({"purchase": "buy", "redemption": "sell"})
    out["date"] = out["date"].dt.strftime("%Y-%m-%d")
    return out.dropna(subset=["date", "scheme_name", "units"]).reset_index(drop=True)


# ---------- helpers ----------

def _find_header_row(raw: pd.DataFrame, must_contain: set[str]) -> int:
    """Locate the row index whose cells contain every label in `must_contain`."""
    for idx, row in raw.iterrows():
        cells = {str(v).strip() for v in row.tolist()}
        if must_contain.issubset(cells):
            return int(idx)
    raise ValueError(
        f"Could not locate header row containing {must_contain} in sheet. "
        "File format may have changed."
    )


def _build_symbol_map(raw_stock_tx: pd.DataFrame) -> dict[str, str]:
    df = raw_stock_tx.dropna(subset=["Stock name", "Symbol"])
    mapping: dict[str, str] = {}
    for name, symbol in zip(df["Stock name"].astype(str), df["Symbol"].astype(str)):
        key = name.strip()
        val = symbol.strip().upper()
        if key and val:
            mapping.setdefault(key, val)
    return mapping


def _numeric_with_commas(s: pd.Series) -> pd.Series:
    return pd.to_numeric(
        s.astype(str).str.replace(",", "", regex=False).str.strip(),
        errors="coerce",
    )
